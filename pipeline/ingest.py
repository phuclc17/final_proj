"""
Tang Bronze: nap file tu cac he thong nguon vao DuckDB (schema `bronze`).

Nguyen tac:
- Giu NGUYEN gia tri goc: moi gia tri luu dang chuoi (bang) hoac JSON goc (payload). Khong ep kieu, khong lam sach.
- Chi them metadata: _source_file, _file_hash, _row_number, _batch_id, _ingested_at (+ cac dong tieu de cua file Excel).
- Nap tang dan, idempotent: file da nap voi cung noi dung thi bo qua; file bi sua thi xoa ban cu va nap lai.
- Moi file duoc ghi vao bronze._ingest_log; moi nguon nap trong 1 transaction.
- Tep nam dung thu muc nhung sai hop dong du lieu cua nguon (pipeline/contracts.py) -> khong nap, status 'rejected'.
- Tep khong thuoc nguon nao (sai ten, sai cho, nguon la) -> khong nap, status 'unrecognized'; KHONG bo qua im lang.

  python -m pipeline.ingest --profile calibration
"""
import argparse
import csv
import hashlib
import io
import json
import re
import time
import unicodedata
import uuid
from datetime import datetime, timezone
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook

from . import dbutil

ROOT = Path(__file__).resolve().parent.parent
META_TYPES = {"_row_number": "INTEGER", "_ingested_at": "TIMESTAMP"}


def slug(name):
    """Ten cot an toan cho SQL: bo dau tieng Viet, ky tu dac biet -> '_' (chi doi TEN cot, khong doi gia tri)."""
    s = str(name).replace("Đ", "D").replace("đ", "d")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^0-9a-zA-Z]+", "_", s).strip("_").lower() or "col"


def file_md5(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def cell_text(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    return str(v)


# ---------------------------------------------------------------- bo doc theo dinh dang file
def read_json_list(path):
    return [{"payload": json.dumps(r, ensure_ascii=False)} for r in json.load(open(path, encoding="utf-8"))]


def read_json_wrapped(path, key="data"):
    doc = json.load(open(path, encoding="utf-8"))
    envelope = json.dumps({k: v for k, v in doc.items() if k != key}, ensure_ascii=False)
    return [{"payload": json.dumps(r, ensure_ascii=False), "_envelope": envelope} for r in doc[key]]


def read_json_object(path):
    return [{"payload": json.dumps(json.load(open(path, encoding="utf-8")), ensure_ascii=False)}]


def read_csv(path):
    """CSV co the den voi nhieu bang ma (UTF-8 co BOM cua Excel, UTF-16 cua phan mem may cham cong cu, Windows-1258)."""
    from .contracts import read_text
    text, _ = read_text(Path(path).read_bytes())
    if text is None:
        raise ValueError("không xác định được bảng mã ký tự")
    reader = csv.reader(io.StringIO(text, newline=""))
    header = [slug(h) for h in next(reader)]
    return [dict(zip(header, row)) for row in reader]


def read_excel(path):
    """Excel dang bang hoac dang phieu: cac dong truoc dong tieu de cot duoc giu lai lam _doc_line_k."""
    wb = load_workbook(path, read_only=True, data_only=True)
    grid = [[cell_text(v) for v in row] for row in wb.active.iter_rows(values_only=True)]
    wb.close()
    header_idx = next(i for i, row in enumerate(grid) if sum(v is not None for v in row) >= 3)
    doc_lines = [" ".join(v for v in row if v) for row in grid[:header_idx] if any(row)]
    header = [slug(h) if h else f"col_{k + 1}" for k, h in enumerate(grid[header_idx])]
    doc = {f"_doc_line_{k + 1}": line for k, line in enumerate(doc_lines)}
    rows = []
    for row in grid[header_idx + 1:]:
        if any(v is not None for v in row):
            rows.append({**dict(zip(header, row)), **doc})
    return rows


def sources(src):
    """(bang bronze, danh sach file, ham doc). Day la 'cau hinh nap' - them nguon moi chi can them 1 dong."""
    yield "pos_receipts", sorted(src.glob("pos/*/*.json")), read_json_list
    yield "app_orders", sorted(src.glob("owned_app/*.json")), read_json_wrapped
    yield "grab_settlement", sorted(src.glob("grabfood/*.csv")), read_csv
    yield "shopee_settlement", sorted(src.glob("shopeefood/*.xlsx")), read_excel
    yield "warehouse_receipts_form", sorted(src.glob("warehouse_receipts/*/NhapKho_*.xlsx")), read_excel
    yield "warehouse_receipts_table", sorted(src.glob("warehouse_receipts/*/receipts_*.xlsx")), read_excel
    yield "stock_counts", sorted(src.glob("stock_counts/*/*.xlsx")), read_excel
    yield "timekeeping_punches", sorted(src.glob("timekeeping/*/*.csv")), read_csv
    yield "weather_daily", sorted(src.glob("external/weather/*.json")), read_json_object
    for entity_dir in sorted(p for p in (src / "master_data").iterdir() if p.is_dir()):
        yield f"master_{entity_dir.name}", sorted(entity_dir.glob("*.csv")), read_csv


# ---------------------------------------------------------------- ghi DuckDB
def table_exists(con, table):
    return con.execute("select count(*) from information_schema.tables where table_schema = 'bronze' and table_name = ?",
                       [table]).fetchone()[0] > 0


def write_rows(con, table, df):
    con.register("incoming", df)
    if not table_exists(con, table):
        cols = ", ".join(f'"{c}" {META_TYPES.get(c, "VARCHAR")}' for c in df.columns)
        con.execute(f"create table bronze.{table} ({cols})")
    else:
        existing = {r[0] for r in con.execute(
            "select column_name from information_schema.columns where table_schema = 'bronze' and table_name = ?",
            [table]).fetchall()}
        for c in df.columns:
            if c not in existing:  # nguon doi schema (vd them cot) -> them cot, khong lam vo pipeline
                con.execute(f'alter table bronze.{table} add column "{c}" VARCHAR')
    con.execute(f"insert into bronze.{table} by name select * from incoming")
    con.unregister("incoming")


IGNORED = {"desktop.ini", "thumbs.db", ".ds_store"}


def stray_files(src, known):
    """Tep trong vung tiep nhan khong thuoc nguon nao da khai bao (bo qua tep he thong, tep khoa cua Excel)."""
    out = []
    for p in sorted(src.rglob("*")):
        if p.is_file() and p not in known and p.name.lower() not in IGNORED and not p.name.startswith(("~$", ".")):
            out.append(p)
    return out


def run(profile):
    src = ROOT / "data" / profile / "source_systems"
    db_path = ROOT / "data" / profile / "warehouse" / "fastfat.duckdb"
    db_path.parent.mkdir(parents=True, exist_ok=True)
    con = dbutil.connect(db_path)
    con.execute("create schema if not exists bronze")
    con.execute("""create table if not exists bronze._ingest_log (
        batch_id varchar, bronze_table varchar, file_path varchar, file_hash varchar, file_bytes bigint,
        row_count integer, status varchar, error varchar, ingested_at timestamp)""")
    loaded = dict(con.execute("""
        select file_path, file_hash from bronze._ingest_log where status in ('loaded', 'reloaded')
        qualify row_number() over (partition by file_path order by ingested_at desc) = 1""").fetchall())
    # tep da bi tu choi / khong nhan dang voi cung noi dung: khong ghi nhat ky lap lai moi luot
    judged = set(con.execute("select file_path, file_hash from bronze._ingest_log "
                             "where status in ('rejected', 'unrecognized')").fetchall())
    # tep cung kich thuoc va khong sua tu lan danh gia gan nhat -> khong doc lai de bam MD5 (cung quy tac voi buoc
    # check cua run_daily). Doc lai ca nghin tep qua thu muc gan tu Windows vao Docker la phan cham nhat cua buoc nap.
    seen = {p: (b, t) for p, b, t in con.execute(
        "select file_path, file_bytes, ingested_at from bronze._ingest_log "
        "where status in ('loaded', 'reloaded', 'rejected', 'unrecognized') "
        "qualify row_number() over (partition by file_path order by ingested_at desc) = 1").fetchall()}

    def unchanged(path, rel):
        old = seen.get(rel)
        if old is None:
            return False
        st = path.stat()
        mtime = datetime.fromtimestamp(st.st_mtime, timezone.utc).replace(tzinfo=None)
        return old[0] == st.st_size and mtime <= old[1]
    from .contracts import check_for_ingest, identify, BY_KEY

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    batch_id = now.strftime("%Y%m%dT%H%M%S") + "-" + uuid.uuid4().hex[:6]
    summary = []
    for table, files, reader in sources(src):
        t0 = time.time()
        rows, log, skipped = [], [], 0
        con.begin()
        try:
            for path in files:
                rel = path.relative_to(src).as_posix()
                if unchanged(path, rel):
                    skipped += 1
                    continue
                h = file_md5(path)
                if loaded.get(rel) == h:
                    skipped += 1
                    continue
                status = "reloaded" if rel in loaded else "loaded"
                if (rel, h) in judged:
                    skipped += 1
                    continue
                errs = check_for_ingest(table, path, src)
                if errs:   # sai hop dong: khong nap, ban cu (neu co) giu nguyen
                    log.append((batch_id, table, rel, h, path.stat().st_size, 0, "rejected", "; ".join(errs)[:500], now))
                    continue
                try:
                    file_rows = reader(path)
                except Exception as e:  # file hong khong duoc lam dung ca lo nap; ban cu (neu co) duoc giu nguyen
                    log.append((batch_id, table, rel, h, path.stat().st_size, 0, "failed", repr(e)[:500], now))
                    continue
                # chi xoa ban cu SAU KHI doc duoc ban moi: ban sua bi hong khong lam mat du lieu dang dung
                if status == "reloaded" and table_exists(con, table):
                    con.execute(f"delete from bronze.{table} where _source_file = ?", [rel])
                for i, r in enumerate(file_rows, start=1):
                    r.update(_source_file=rel, _file_hash=h, _row_number=i)
                rows.extend(file_rows)
                log.append((batch_id, table, rel, h, path.stat().st_size, len(file_rows), status, None, now))
            if rows:
                df = pd.DataFrame(rows)
                for c in df.columns:
                    if c != "_row_number":
                        df[c] = df[c].astype("string")
                df["_batch_id"] = batch_id
                df["_ingested_at"] = now
                write_rows(con, table, df)
            if log:
                con.executemany("insert into bronze._ingest_log values (?, ?, ?, ?, ?, ?, ?, ?, ?)", log)
            con.commit()
        except Exception:
            con.rollback()
            raise
        failed = sum(1 for x in log if x[6] in ("failed", "rejected"))
        summary.append((table, len(log) - failed, skipped, failed, len(rows), round(time.time() - t0, 1)))

    # tep khong thuoc nguon nao: ghi nhan de canh bao, goi y vi tri dung neu noi dung nhan dang duoc
    known = {p for _, files, _ in sources(src) for p in files}
    stray = []
    for path in stray_files(src, known):
        rel = path.relative_to(src).as_posix()
        if unchanged(path, rel):
            continue
        h = file_md5(path)
        if (rel, h) in judged:
            continue
        r = identify(path.name, path.read_bytes(), src) if path.stat().st_size < 50 << 20 else None
        if r is not None and r.source and r.target:
            msg = (f"Không đúng tên hoặc vị trí. Nội dung là {BY_KEY[r.source]['system']}; vị trí đúng: {r.target}. "
                   "Dùng kênh tải tệp lên để hệ thống tự xếp đúng chỗ.")
        else:
            msg = "Không thuộc nguồn nào. " + ("; ".join(r.errors) if r is not None else "Tệp quá lớn.")
        stray.append((batch_id, None, rel, h, path.stat().st_size, 0, "unrecognized", msg[:500], now))
    if stray:
        con.executemany("insert into bronze._ingest_log values (?, ?, ?, ?, ?, ?, ?, ?, ?)", stray)

    print(f"Lo nap {batch_id} -> {db_path.relative_to(ROOT)}")
    print(pd.DataFrame(summary, columns=["bang_bronze", "file_moi", "file_bo_qua", "file_loi", "dong_moi", "giay"])
          .to_string(index=False))
    tot = pd.DataFrame(summary, columns=["t", "moi", "bo_qua", "loi", "dong", "s"]).sum(numeric_only=True)
    print(f"Tong: {int(tot.moi)} file moi, {int(tot.dong)} dong moi, {int(tot.loi)} file loi"
          + (f", {len(stray)} file khong nhan dang" if stray else ""))
    con.close()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--profile", default="calibration", choices=["calibration", "holdout", "live"])
    run(ap.parse_args().profile)
