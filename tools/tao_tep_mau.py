"""Sinh bo tep mau cho kenh "Tai tep len": moi tep la mot tinh huong hay gap ngoai thuc te khi nguoi gui tep thu cong.
Lay tu anh chup demo cua moi truong van hanh (du lieu den 22/09/2024), nen dung duoc ca truoc va sau khi demo sinh them
ngay moi. Ket qua mong doi cua tung tep ghi trong samples/tep_mau/DOC_TRUOC.txt.

    .venv/Scripts/python tools/tao_tep_mau.py
"""
import csv
import io
import json
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
SNAP = ROOT / "data" / "_demo_snapshot" / "live" / "source_systems"
SRC = SNAP if SNAP.exists() else ROOT / "data" / "live" / "source_systems"
OUT = ROOT / "samples" / "tep_mau"

CASES = []


def case(name, expect, why):
    CASES.append((name, expect, why))
    return OUT / name


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    # 1. phieu kiem ke gui lai ban sua: sua so dem cua mot nguyen lieu
    wb = load_workbook(SRC / "stock_counts" / "CN03" / "KiemKe_CN03_20240922.xlsx")
    ws = wb.active
    for row in ws.iter_rows():
        if row[2].value == "Thịt gà tươi":
            row[4].value = round(float(row[4].value) + 1.5, 1)
            row[5].value = "đếm lại, bổ sung 1,5 kg trong tủ đông phụ"
            break
    wb.save(case("01_KiemKe_CN03_22-09_ban_sua.xlsx", "thay bản cũ",
                 "Phiếu kiểm kê CN03 ngày 22/09 gửi lại sau khi đếm lại một nguyên liệu. Cùng chi nhánh, cùng ngày, "
                 "nội dung khác: hệ thống coi là bản sửa, nạp lại thay bản cũ."))

    # 2. bao cao ShopeeFood gui trung nhung dat ten khac
    shutil.copy(SRC / "shopeefood" / "ShopeeFood_DoiSoat_20240916_20240922.xlsx",
                case("02_bao_cao_shopee_tuan_38.xlsx", "trùng, bỏ qua",
                     "Kế toán gửi lại báo cáo ShopeeFood tuần 16–22/09 dưới tên khác. Hệ thống nhận ra theo nội dung, "
                     "không theo tên tệp, nên không nạp hai lần."))

    # 3. bao cao GrabFood luu nham sang Excel
    df = pd.read_csv(SRC / "grabfood" / "GrabFood_Settlement_20240916_20240922.csv", dtype=str)
    df.to_excel(case("03_GrabFood_tuan_38.xlsx", "từ chối",
                     "Báo cáo GrabFood bị mở bằng Excel rồi lưu lại dạng .xlsx. GrabFood xuất CSV; hệ thống từ chối và "
                     "chỉ rõ cần gửi tệp CSV gốc."), index=False)

    # 4. quan ly bo sung luot ra bi quen, tep xuat tu phan mem may cham cong cu (UTF-16)
    src_tk = SRC / "timekeeping" / "BR02" / "punch_BR02_20240922.csv"
    rows = list(csv.reader(open(src_tk, encoding="utf-8")))
    rows.append(["14", "BR02", "CHECK_OUT", "2024-09-22 22:06:10", "DEV-TK-02", "manual"])
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows(rows)
    case("04_chamcong_BR02_22-09_bo_sung_luot_ra.csv", "thay bản cũ",
         "Quản lý CN02 bổ sung lượt ra mà nhân viên 14 quên quẹt tối 22/09, xuất từ phần mềm máy chấm công cũ "
         "(bảng mã UTF-16). Hệ thống chuyển về UTF-8 rồi nạp lại thay bản cũ.").write_bytes(buf.getvalue().encode("utf-16"))

    # 5. tep cham cong thieu cot
    df = pd.read_csv(SRC / "timekeeping" / "BR04" / "punch_BR04_20240922.csv", dtype=str).drop(columns=["event_type"])
    df.to_csv(case("05_chamcong_BR04_thieu_cot.csv", "từ chối",
                   "Tệp chấm công bị xóa mất cột event_type khi chỉnh sửa bằng tay: không biết lượt nào là vào, lượt nào "
                   "là ra. Hệ thống từ chối, không đoán."), index=False)

    # 6. so nhap kho mau B dat ten khong co ma chi nhanh
    shutil.copy(SRC / "warehouse_receipts" / "CN06" / "receipts_CN06_20240921.xlsx",
                case("06_so_nhap_kho_tuan_nay.xlsx", "từ chối",
                     "Sổ nhập kho mẫu B không ghi chi nhánh trong tệp, và tên tệp cũng không có mã CNxx. Hệ thống không "
                     "biết là của chi nhánh nào nên từ chối, yêu cầu đặt lại tên."))

    # 7. tep la khong thuoc nguon nao
    wb = Workbook()
    ws = wb.active
    ws.append(["BẢNG TỔNG HỢP LƯƠNG THÁNG 9/2024"])
    ws.append([])
    ws.append(["Mã NV", "Họ tên", "Chi nhánh", "Ngày công", "Lương"])
    ws.append(["NV0014", "Nguyễn Văn A", "CN02", 24, 6200000])
    wb.save(case("07_bang_luong_thang_9.xlsx", "từ chối",
                 "Tệp không thuộc nguồn dữ liệu nào của nền tảng (bảng lương nội bộ). Hệ thống từ chối là không nhận dạng "
                 "được, không bỏ qua im lặng."))

    # 8. POS cua mot may chua co trong bang anh xa (chi nhanh moi chua khai bao)
    recs = json.load(open(SRC / "pos" / "POS-05" / "pos_export_2024-09-22.json", encoding="utf-8"))[:3]
    for r in recs:
        for k in ("receipt_no", "terminal_id", "pos_device_id", "shift_id"):
            if r.get(k):
                r[k] = str(r[k]).replace("POS-05", "POS-12")
    case("08_pos_may_moi_POS-12.json", "từ chối",
         "Hóa đơn từ máy POS-12, của chi nhánh mới chưa được khai báo trong bảng ánh xạ mã. Hệ thống từ chối cả tệp "
         "thay vì nạp doanh thu không gắn được chi nhánh.").write_text(json.dumps(recs, ensure_ascii=False, indent=1),
                                                                         encoding="utf-8")

    # 9. phieu nhap mua gap bo sung, den muon (ngay Chu nhat chi nhanh khong co phieu nao)
    wb = Workbook()
    ws = wb.active
    for row in (["PHIẾU NHẬP KHO"], ["Chi nhánh: Fast&Fat 250 Ly Chinh Thang"], ["Ngày nhập: 22/09/2024"], [],
                ["STT", "Tên nguyên liệu", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền", "Hạn sử dụng", "Nhà cung cấp"],
                [1, "Nước suối", "chai", 24, 2600, 62400, "22/09/2025", "Cửa hàng tiện lợi gần chi nhánh"],
                [2, "Pepsi lon", "lon", 24, 4300, 103200, "22/09/2025", "Cửa hàng tiện lợi gần chi nhánh"],
                ["", "Tổng cộng", "", "", "", 165600], [], ["Người lập phiếu: Quản lý CN01"]):
        ws.append(row)
    wb.save(case("09_NhapKho_CN01_22-09_mua_gap.xlsx", "đã nhận",
                 "Phiếu mua gấp nước uống ngày Chủ nhật 22/09 ở CN01, quản lý gửi bổ sung muộn. Hợp lệ, chưa có phiếu "
                 "nào của CN01 ngày đó: hệ thống nhận, nạp và tính lại tồn kho của kỳ."))

    lines = ["Bộ tệp mẫu để thử kênh Tải tệp lên (môi trường vận hành). Kết quả mong đợi khi gửi lần đầu:", ""]
    lines += [f"{n}\n    Kết quả mong đợi: {e}\n    {w}\n" for n, e, w in CASES]
    lines += ["Gửi lại cùng một tệp lần thứ hai thì kết quả là 'trùng, bỏ qua' (hoặc 'từ chối' với tệp sai).",
              "Muốn thử lại từ đầu: .venv\\Scripts\\python tools\\chuan_bi_demo.py --khoi-phuc"]
    (OUT / "DOC_TRUOC.txt").write_text("\n".join(lines), encoding="utf-8")
    print(f"Đã tạo {len(CASES)} tệp mẫu ở {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
