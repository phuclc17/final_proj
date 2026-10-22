"""Ghi file theo dung dinh dang moi he thong nguon xuat ra: JSON (POS, App), CSV (Grab, cham cong), Excel (Shopee, kho)."""
import csv
import json
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

from . import config as C
from . import views
from .reference import INGREDIENT_VN, MENU_VN, UNIT_VN, app_product_ref
from .world import pos_code


class Exporter:
    def __init__(self, profile_dir, api_store=None):
        self.src = profile_dir / "source_systems"
        self.api = api_store   # None: moi he thong xuat tep vao vung tiep nhan (tap phat trien, tap kiem tra)
        self.truth = profile_dir / "ground_truth"
        self.src.mkdir(parents=True, exist_ok=True)
        self.truth.mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------ tien ich
    def _path(self, *parts, api=False):
        p = (self.api if api and self.api else self.src).joinpath(*parts)
        p.parent.mkdir(parents=True, exist_ok=True)
        return p

    @staticmethod
    def _write_csv(path, rows, fields, mode="w"):
        new = mode == "w" or not path.exists()
        with open(path, mode, newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            if new:
                w.writeheader()
            w.writerows(rows)

    @staticmethod
    def _write_json(path, obj):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))

    # ------------------------------------------------------------ 1 ngay
    def write_day(self, world, out):
        d = out["date"]
        ymd = d.strftime("%Y%m%d")
        ref = world.ref
        for entity in dict.fromkeys(out["master"]):
            self._master(world, entity, ymd)

        all_orders = []
        for b, orders in out["orders"].items():
            prefix = ref.branches[b]["names"]["pos_terminal_prefix"]
            self._write_json(self._path("pos", prefix, f"pos_export_{d.isoformat()}.json",
                                        api=b not in C.OLD_POS_BRANCHES),
                             views.pos_receipts(world, b, d, orders))
            all_orders.extend(orders)
        app = views.app_orders(world, all_orders)
        self._write_json(self._path("owned_app", f"orders_{ymd}.json", api=True),
                         {"export_generated_at": f"{(d + timedelta(days=1)).isoformat()}T01:00:00Z",
                          "page": 1, "page_size": len(app), "total": len(app), "data": app})

        for b, rows in out["punches"].items():
            self._timekeeping(world, b, ymd, rows)
        for b, rows in out["receipts"].items():
            if rows:
                self._receipt_excel(world, b, d, rows)
        for b, rows in out["counts"].items():
            self._count_excel(world, b, d, rows)
        if out["settlement"]:
            self._grab_csv(world, out["settlement"])
            self._shopee_excel(world, out["settlement"])
        self._write_json(self._path("external", "weather", f"weather_hcm_{ymd}.json"), out["weather"])

        truth_fields = list(out["truth"][0].keys())
        self._write_csv(self.truth / "daily_branch_truth.csv", out["truth"], truth_fields, mode="a")
        if out["counts"]:
            rows = [{"branch_id": b, "count_date": d.isoformat(), "ingredient_id": r["ingredient_id"],
                     "true_stock": r["true_stock"]} for b, rs in out["counts"].items() for r in rs]
            self._write_csv(self.truth / "stock_truth.csv", rows, list(rows[0].keys()), mode="a")
        if out["settlement"]:
            s = out["settlement"]
            rows = [{"week_start": s["week_start"].isoformat(), "platform": p, "branch_id": r["branch_id"],
                     "commission_rate_applied": r["commission_rate_applied"]}
                    for p, rs in s["rows"].items() for r in rs]
            self._write_csv(self.truth / "commission_truth.csv", rows, list(rows[0].keys()), mode="a")

    def write_scenarios(self, world):
        items = [{k: (v.isoformat() if hasattr(v, "isoformat") else v) for k, v in s.items()} for s in world.book.items]
        self._write_json(self.truth / "scenarios.json", {"profile": world.book.profile, "seed": world.book.seed,
                                                         "scenarios": items})

    # ------------------------------------------------------------ master data (xuat tu he thong HQ)
    def _master(self, world, entity, ymd):
        ref = world.ref
        if entity == "branches":
            rows = [{"branch_id": b, "branch_name": r["branch_name"], "district": r["district"],
                     "latitude": r["latitude"], "longitude": r["longitude"], "open_date": r["open_date"],
                     "store_class": C.STORE_CLASS_BY_TIER[r["tier"]]} for b, r in ref.branches.items()]
        elif entity == "menu_items":
            rows = [{"item_id": i, "item_name": m["item_name"], "category": m["category"],
                     "price": world.prices[i], "standard_cost": m["standard_cost"]} for i, m in ref.menu.items()]
        elif entity == "app_catalog":
            rows = [{"product_ref": app_product_ref(i), "product_name": MENU_VN[i],
                     "category_ref": f"CAT-{m['category'][:3].upper()}", "price_vnd": world.prices[i]}
                    for i, m in ref.menu.items()]
        elif entity == "ingredients":
            rows = [{"ingredient_id": i, "ingredient_code": f"NL{i:02d}", "ingredient_name": g["ingredient_name"],
                     "unit": g["unit"], "category": g["category"], "shelf_life_days": g["shelf_life_days"],
                     "standard_unit_cost": int(g["unit_cost"])} for i, g in ref.ingredients.items()]
        elif entity == "recipes":
            rows = ref.recipes_raw
        elif entity == "combo_components":
            rows = ref.combos_raw
        elif entity == "staff":
            rows = [{"staff_id": s["staff_id"], "staff_name": s["staff_name"], "branch_id": s["branch_id"],
                     "role": s["role"], "hourly_wage": s["hourly_wage"], "hire_date": s["hire_date"]}
                    for s in world.staff.values()]
        elif entity == "marketplace_contracts":
            rows = [{"branch_id": b, "platform": p, "commission_rate": C.COMMISSION_CONTRACT[p],
                     "effective_from": C.START_DATE.isoformat()} for b in ref.branches for p in C.MARKETPLACES]
        else:
            raise ValueError(entity)
        self._write_csv(self._path("master_data", entity, f"{entity}_{ymd}.csv"), rows, list(rows[0].keys()))

    # ------------------------------------------------------------ cham cong (may van tay xuat CSV)
    def _timekeeping(self, world, b, ymd, rows):
        code = world.ref.branches[b]["names"]["timekeeping_branch_code"]
        rng = world.rng
        out = [{"employee_code": str(r["staff_id"]), "branch_code": code, "event_type": r["event_type"],
                "event_time": r["event_time"].strftime("%Y-%m-%d %H:%M:%S"), "device_id": f"DEV-TK-{b:02d}",
                "verification_method": ["fingerprint", "fingerprint", "face", "pin"][int(rng.integers(4))]}
               for r in rows]
        self._write_csv(self._path("timekeeping", code, f"punch_{code}_{ymd}.csv"), out,
                        ["employee_code", "branch_code", "event_type", "event_time", "device_id", "verification_method"])

    # ------------------------------------------------------------ so nhap kho (Excel quan ly tu ghi)
    def _receipt_excel(self, world, b, d, rows):
        template, lines = views.receipt_rows(world, b, rows)
        folder = f"CN{b:02d}"
        wb = Workbook()
        ws = wb.active
        if template == "A":
            ws.title = "PhieuNhap"
            ws.append(["PHIẾU NHẬP KHO"])
            ws["A1"].font = Font(bold=True, size=14)
            ws.merge_cells("A1:H1")
            ws.append([f"Chi nhánh: {world.ref.branches[b]['branch_name']}"])
            ws.append([f"Ngày nhập: {d.strftime('%d/%m/%Y')}"])
            ws.append([])
            ws.append(["STT", "Tên nguyên liệu", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền", "Hạn sử dụng", "Nhà cung cấp"])
            total = 0
            for k, r in enumerate(lines, start=1):
                amount = round(r["quantity"] * r["unit_price"])
                total += amount
                ws.append([k, r["name"], r["unit"], r["quantity"], r["unit_price"], amount,
                           r["expiry"].strftime("%d/%m/%Y"), r["supplier"]])
            ws.append([None, "Tổng cộng", None, None, None, total])
            ws.append([])
            ws.append([f"Người lập phiếu: Quản lý CN{b:02d}"])
            name = f"NhapKho_CN{b:02d}_{d.strftime('%Y%m%d')}.xlsx"
        else:
            ws.title = "receipts"
            ws.append(["date", "item", "quantity", "unit_price", "supplier"])
            for r in lines:
                ws.append([r["date"].isoformat(), r["name"], f"{r['quantity']}{r['unit']}", r["unit_price"], r["supplier"]])
            name = f"receipts_CN{b:02d}_{d.strftime('%Y%m%d')}.xlsx"
        wb.save(self._path("warehouse_receipts", folder, name))

    # ------------------------------------------------------------ phieu kiem ke (Excel)
    def _count_excel(self, world, b, d, rows):
        ref = world.ref
        managers = [s["staff_name"] for s in world.staff.values() if s["branch_id"] == b and s["role"] == "ShiftManager"]
        wb = Workbook()
        ws = wb.active
        ws.title = "KiemKe"
        ws.append(["PHIẾU KIỂM KÊ KHO"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:F1")
        ws.append([f"Chi nhánh: {ref.branches[b]['branch_name']}"])
        ws.append([f"Ngày kiểm kê: {d.strftime('%d/%m/%Y')}"])
        ws.append([f"Người kiểm: {managers[0] if managers else ''}"])
        ws.append([])
        ws.append(["STT", "Mã NL", "Tên nguyên liệu", "ĐVT", "Tồn thực tế", "Ghi chú"])
        for k, r in enumerate(rows, start=1):
            ing = r["ingredient_id"]
            ws.append([k, f"NL{ing:02d}", INGREDIENT_VN[ing][0], UNIT_VN[ref.ingredients[ing]["unit"]],
                       r["counted"], None if r["counted"] is not None else "chưa đếm"])
        wb.save(self._path("stock_counts", f"CN{b:02d}", f"KiemKe_CN{b:02d}_{d.strftime('%Y%m%d')}.xlsx"))

    # ------------------------------------------------------------ doi soat san
    def _grab_csv(self, world, s):
        ref = world.ref
        start, end = s["week_start"], s["week_end"]
        rows = [{"Merchant Store Name": ref.branches[r["branch_id"]]["names"]["grabfood_settlement_name"],
                 "Week": f"{start.strftime('%d/%m/%Y')} - {end.strftime('%d/%m/%Y')}",
                 "Report Generated Date": (end + timedelta(days=2)).strftime("%d %b %Y"),
                 "Gross Sales": r["gross_sales"], "Total Orders": r["orders_completed"],
                 "Cancelled Orders": r["orders_cancelled"], "Total Commission": r["commission"],
                 "Ads Spend": r["ads_spend"], "Promotion Cost Borne by Merchant": r["merchant_promo"],
                 "Adjustment Amount": r["adjustment"], "Average Prep Time (min)": r["avg_prep_min"],
                 "Store Rating": r["store_rating"], "Net Payout": r["net_payout"]}
                for r in s["rows"]["GRABFOOD"]]
        name = f"GrabFood_Settlement_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.csv"
        self._write_csv(self._path("grabfood", name), rows, list(rows[0].keys()))

    def _shopee_excel(self, world, s):
        ref = world.ref
        start, end = s["week_start"], s["week_end"]
        wb = Workbook()
        ws = wb.active
        ws.title = "DoiSoat"
        ws.append(["BÁO CÁO ĐỐI SOÁT ShopeeFood"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Kỳ đối soát: {start.strftime('%d/%m/%Y')} - {end.strftime('%d/%m/%Y')}"])
        ws.append([])
        ws.append(["Mã cửa hàng", "Số đơn hoàn thành", "Số đơn hủy", "Doanh thu món (VNĐ)", "Phí dịch vụ (VNĐ)",
                   "Khuyến mãi do quán tài trợ (VNĐ)", "Phí quảng cáo (VNĐ)", "Điều chỉnh (VNĐ)", "Thực nhận (VNĐ)",
                   "Thời gian chuẩn bị TB (phút)", "Đánh giá"])
        for r in s["rows"]["SHOPEEFOOD"]:
            ws.append([ref.branches[r["branch_id"]]["names"]["shopeefood_settlement_name"], r["orders_completed"],
                       r["orders_cancelled"], r["gross_sales"], r["commission"], r["merchant_promo"], r["ads_spend"],
                       r["adjustment"], r["net_payout"], r["avg_prep_min"], r["store_rating"]])
        for row in ws.iter_rows(min_row=5, min_col=4, max_col=9):
            for cell in row:
                cell.number_format = "#,##0"
        name = f"ShopeeFood_DoiSoat_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.xlsx"
        wb.save(self._path("shopeefood", name))
