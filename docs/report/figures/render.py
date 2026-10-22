"""
Dung hinh cho bao cao: moi hinh co 1 tep nguon trong figures/src/ (Graphviz .dot hoac trang .html),
xuat ra figures/<ten>.png do phan giai gap doi bang trinh duyet khong giao dien (Playwright).
Hinh ERD duoc SINH TU DONG tu metadata dbt (manifest.json), hinh du lieu tho sinh tu tep nguon that cua tap
phat trien, nen luon khop voi he thong.

  .venv/Scripts/python docs/report/figures/render.py            # dung tat ca
  .venv/Scripts/python docs/report/figures/render.py h3_5_erd   # dung 1 hinh
"""
import json
import re
import sys
from pathlib import Path

from playwright.sync_api import sync_playwright

HERE = Path(__file__).resolve().parent
SRC, ROOT = HERE / "src", HERE.parents[2]
DOT_PAGE = """<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="_style.css">
<script src="../vendor/viz-standalone.js"></script></head><body><div id="fig"></div>
<script>Viz.instance().then(v => {{ const svg = v.renderSVGElement({dot}); document.getElementById('fig').appendChild(svg);
document.body.dataset.ready = 1; }});</script></body></html>"""


# ---------------------------------------------------------------- hinh sinh tu dong
def gen_erd():
    """So do hinh sao tu manifest dbt: bang su kien o giua, bang chieu xung quanh, cot khoa that cua tung bang."""
    m = json.loads((ROOT / "dbt" / "target_calibration" / "manifest.json").read_text(encoding="utf-8"))
    gold = {n["name"]: n for n in m["nodes"].values() if n["resource_type"] == "model" and n["schema"] == "gold"
            and n["name"].startswith(("dim_", "fact_"))}
    rel = set()
    for n in m["nodes"].values():
        meta = n.get("test_metadata") or {}
        if meta.get("name") != "relationships":
            continue
        src = (n.get("attached_node") or "").split(".")[-1]
        dst = re.search(r"ref\('([^']+)'\)", meta["kwargs"].get("to", ""))
        if src in gold and dst and dst.group(1) in gold:
            rel.add((src, dst.group(1), meta["kwargs"].get("column_name"), meta["kwargs"].get("field")))
    fk = {(a, c) for a, b, c, f in rel}
    pk = {(b, f) for a, b, c, f in rel}
    grain = {"fact_order": "order_key", "fact_order_line": "order_line_key", "fact_labor_shift": "shift_key",
             "fact_marketplace_settlement": "settlement_key", "fact_inventory_receipt": "receipt_line_id",
             "fact_stock_count": "count_key", "dim_menu_item": "menu_item_sk", "dim_employee": "employee_sk",
             "dim_branch": "branch_id", "dim_date": "date_day", "dim_channel": "channel",
             "dim_ingredient": "ingredient_id", "dim_customer": "customer_id"}
    extra = {"fact_order": ["net_sales", "order_status"], "fact_order_line": ["qty", "cogs_standard"],
             "fact_labor_shift": ["hours_paid", "labor_cost"], "fact_marketplace_settlement": ["commission_rate_gap"],
             "fact_inventory_receipt": ["quantity", "unit_price"], "fact_stock_count": ["counted_qty"],
             "dim_menu_item": ["price", "valid_from", "valid_to"], "dim_employee": ["role", "valid_from", "valid_to"],
             "dim_branch": ["store_class"], "dim_date": ["is_holiday", "is_rain"], "dim_channel": ["channel_group"],
             "dim_ingredient": ["is_fresh"], "dim_customer": ["first_order_date"]}
    # khoa noi khong co kiem thu relationships nhung la noi tu nhien theo ma tran bus
    natural = [("fact_order", "dim_date", "date_day"), ("fact_order_line", "dim_date", "date_day"),
               ("fact_order_line", "dim_branch", "branch_id"), ("fact_labor_shift", "dim_branch", "branch_id"),
               ("fact_labor_shift", "dim_date", "date_day"), ("fact_marketplace_settlement", "dim_branch", "branch_id"),
               ("fact_marketplace_settlement", "dim_channel", "channel"), ("fact_inventory_receipt", "dim_branch", "branch_id"),
               ("fact_inventory_receipt", "dim_date", "date_day"), ("fact_stock_count", "dim_branch", "branch_id"),
               ("fact_stock_count", "dim_ingredient", "ingredient_id"), ("fact_stock_count", "dim_date", "date_day"),
               ("fact_order_line", "fact_order", "order_key")]

    # Chom sao su kien: dim_date, dim_branch noi voi ca 6 bang su kien -> neu ve het, duong cat cheo khap hinh.
    # Cach trinh bay thuong dung: van liet ke khoa ngoai (to mau xanh) trong tung bang, khong ve duong cua 2 chieu
    # dung chung, chi ve quan he rieng; duong noi tu dung dong khoa ngoai toi dong khoa chinh, ky hieu chan qua.
    conformed = {"dim_date", "dim_branch"}
    grid = [["dim_customer", "fact_order", "fact_order_line"],
            ["dim_channel", "dim_employee", "dim_menu_item"],
            ["fact_marketplace_settlement", "fact_labor_shift", "@conformed"],
            ["fact_inventory_receipt", "dim_ingredient", "fact_stock_count"]]

    target = {(a, c): b for a, b, c, f in rel} | {(a, c): b for a, b, c in natural}
    pos = {n: (r, k) for r, row in enumerate(grid) for k, n in enumerate(row)}

    def table(name):
        is_fact = name.startswith("fact_")
        # khoa ngoai xep theo vi tri bang chieu dich tren luoi (tren -> duoi) de cac duong noi khong cat nhau;
        # khoa toi chieu dung chung xep cuoi
        fks = {c for (t, c) in fk if t == name} | {c for a, b, c in natural if a == name}
        keys = [grain[name]] + sorted(fks - {grain[name]},
                                      key=lambda c: (target[(name, c)] in conformed, pos.get(target[(name, c)], (9, 9)), c))
        rows = [f'<div class="h">{name}</div>']
        for c in keys:
            tag = "PK" if c == grain[name] else "FK"
            cls = "k cf" if tag == "FK" and c in ("date_day", "branch_id") else "k"
            rows.append(f'<div class="{cls}" id="{name}__{c}"><b>{tag}</b> {c}</div>')
        rows += [f'<div class="x">{c}</div>' for c in extra.get(name, [])]
        return f'<div class="t {"f" if is_fact else "d"}" id="{name}">{"".join(rows)}</div>'

    cells = []
    for row in grid:
        for n in row:
            if n == "@conformed":
                cells.append('<div class="cg">' + table("dim_date") + table("dim_branch") +
                             '<div class="nt">Chiều dùng chung: nối với cả 6 bảng sự kiện qua các khóa màu xanh; '
                             'không vẽ đường để hình dễ đọc.</div></div>')
            else:
                cells.append(table(n))
    # (bang su kien, cot khoa ngoai, bang chieu, cot khoa chinh, net dut?)
    links = [(a, c, b, f, False) for a, b, c, f in sorted(rel) if b not in conformed]
    have = {(a, b) for a, b, c, f in rel}
    links += [(a, c, b, c, True) for a, b, c in natural if b not in conformed and (a, b) not in have]
    page = ERD_PAGE.replace("@@CELLS@@", "\n".join(cells)).replace("@@LINKS@@", json.dumps(links))
    (SRC / "h3_5_erd.html").write_text(page, encoding="utf-8")


ERD_PAGE = """<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="_style.css">
<style>
#fig { width: 800px; padding: 4px 26px; box-sizing: border-box; position: relative; }
.g { display: grid; grid-template-columns: 210px 236px 222px; column-gap: 27px; row-gap: 26px; align-items: start; }
.t { border: 1.3px solid #9aa5b1; font-family: Arial, sans-serif; font-size: 13px; background: #fff; }
.t .h { color: #fff; font-weight: 700; padding: 3px 7px; }
.t.f { background: #fff8dc; } .t.f .h { background: #e0b100; }
.t.d { background: #f3f7fb; } .t.d .h { background: #5b7fa6; }
.t .k, .t .x { padding: 2px 7px; line-height: 17px; white-space: nowrap; }
.t .x { color: #52606d; }
.t .cf { color: #3b6fb6; } .t .cf b { color: #3b6fb6; }
.cg { display: grid; grid-template-columns: 1fr 1fr; gap: 8px; }
.cg .nt { grid-column: 1 / -1; font-size: 12px; color: #3e4c59; line-height: 1.35; border: 1px dashed #9aa5b1;
          border-radius: 6px; padding: 4px 7px; background: #fff; }
.lg { margin-top: 14px; font-size: 12px; color: #3e4c59; }
</style></head><body><div id="fig"><div class="g">
@@CELLS@@
</div>
<div class="lg">Nét liền: khóa ngoại có kiểm thử relationships trong dbt · Nét đứt: khóa nối theo ma trận bus ·
  Chân quạ ở phía bảng sự kiện (nhiều), gạch ở phía bảng chiều (một)</div>
</div>
<script>
const fig = document.getElementById("fig"), F = fig.getBoundingClientRect();
const svg = document.createElementNS("http://www.w3.org/2000/svg", "svg");
svg.setAttribute("width", fig.scrollWidth); svg.setAttribute("height", fig.scrollHeight);
svg.style.cssText = "position:absolute;left:0;top:0;pointer-events:none;overflow:visible";
fig.appendChild(svg);
const rect = id => document.getElementById(id).getBoundingClientRect();
function path(d, dash) {
  const p = document.createElementNS("http://www.w3.org/2000/svg", "path");
  p.setAttribute("d", d); p.setAttribute("fill", "none"); p.setAttribute("stroke", "#6b7785");
  p.setAttribute("stroke-width", 1.3); if (dash) p.setAttribute("stroke-dasharray", "5 4"); svg.appendChild(p);
}
// many: chan qua; one: gach doc. s = +1 neu duong di ra ben phai cua bang
function many(x, y, s) { path(`M${x + s * 11},${y} L${x},${y - 6} M${x + s * 11},${y} L${x},${y} M${x + s * 11},${y} L${x},${y + 6}`); }
function one(x, y, s) { path(`M${x + s * 7},${y - 6} L${x + s * 7},${y + 6}`); }
for (const [ft, fc, dt, dc, dash] of @@LINKS@@) {
  const a = rect(ft + "__" + fc), b = rect(dt + "__" + dc), ta = rect(ft), tb = rect(dt);
  const ya = a.top + a.height / 2 - F.top, yb = b.top + b.height / 2 - F.top;
  let xa, xb, sa, sb, d;
  if (Math.abs(ta.left - tb.left) < 30) {            // cung cot: vong ra mot ben
    const right = !["fact_marketplace_settlement", "dim_channel"].includes(ft) && !["dim_channel"].includes(dt);
    sa = sb = right ? 1 : -1;
    xa = (right ? ta.right : ta.left) - F.left; xb = (right ? tb.right : tb.left) - F.left;
    const xo = (right ? Math.max(xa, xb) : Math.min(xa, xb)) + sa * 22;
    d = `M${xa},${ya} C${xo},${ya} ${xo},${yb} ${xb},${yb}`;
  } else {                                            // khac cot: noi canh doi dien
    const toRight = tb.left > ta.left;
    sa = toRight ? 1 : -1; sb = -sa;
    xa = (toRight ? ta.right : ta.left) - F.left; xb = (toRight ? tb.left : tb.right) - F.left;
    const xm = (xa + xb) / 2;
    d = `M${xa},${ya} C${xm},${ya} ${xm},${yb} ${xb},${yb}`;
  }
  path(d, dash); many(xa, ya, sa); one(xb, yb, sb);
}
</script></body></html>"""


def gen_raw():
    """Hinh du lieu tho: dung lai bang tinh, CSV, JSON tu TEP NGUON THAT cua tap phat trien (khong go tay so lieu)."""
    import csv
    import html
    import openpyxl

    src = ROOT / "data" / "calibration" / "source_systems"
    esc = html.escape

    def letters(n):
        return [chr(65 + i) for i in range(n)]

    def sheet(path, rows, cols, hl=lambda r, c, v: False, more=False):
        """Bang tinh nhu Excel: tieu de cot A, B, C..., so dong, o gop, chu tran sang o trong ben phai."""
        ws = openpyxl.load_workbook(path).active
        merged = {(m.min_row, m.min_col): m.max_col - m.min_col + 1 for m in ws.merged_cells.ranges}
        head = "".join(f"<th>{c}</th>" for c in letters(cols)) + ("<th>…</th>" if more else "")
        out = [f'<table class="xl"><tr><th></th>{head}</tr>']
        for r in range(1, rows + 1):
            vals = [ws.cell(r, c).value for c in range(1, cols + 1)]
            tds, c = [], 1
            while c <= cols:
                v = vals[c - 1]
                span = merged.get((r, c), 1)
                if v is not None and span == 1 and c == 1 and all(x is None for x in vals[1:]):
                    span = cols  # chu tran sang cac o trong ben phai nhu Excel hien thi
                cls = ["num"] if isinstance(v, (int, float)) and not isinstance(v, bool) else []
                if (r, c) in merged:
                    cls.append("mg")
                if v is not None and hl(r, c, v):
                    cls.append("hl")
                text = "" if v is None else esc(str(v))
                tds.append(f'<td class="{" ".join(cls)}" colspan="{min(span, cols - c + 1)}">{text}</td>')
                c += span
            tail = "<td></td>" if more else ""
            out.append(f"<tr><th>{r}</th>{''.join(tds)}{tail}</tr>")
        return "".join(out) + "</table>"

    def js(v, ind, hl, key=None):
        pad = "  " * ind
        if isinstance(v, dict):
            items = [f'{pad}  <span class="k">"{k}"</span>: {js(x, ind + 1, hl, k)}' if k != "…" else f"{pad}  …"
                     for k, x in v.items()]
            return "{\n" + ",\n".join(items) + f"\n{pad}}}"
        if isinstance(v, list):
            return "[\n" + ",\n".join(f"{pad}  " + (js(x, ind + 1, hl) if x != "…" else "…") for x in v) + f"\n{pad}]"
        s = esc(json.dumps(v, ensure_ascii=False))
        cls = "s" if isinstance(v, str) else "n"
        return f'<span class="{cls}{" hl" if key in hl else ""}">{s}</span>'

    def receipt(path):
        r = json.loads(path.read_text(encoding="utf-8"))[0]
        tot = "total_paid" if "total_paid" in r else "grand_total"
        hl = {"txn_time", "subtotal_amount", "pos_software_version", tot}
        lines = []
        for k in ("receipt_no", "cashier_code", "txn_time", "items", "subtotal_amount", "…", "pos_software_version", tot):
            if k == "…":
                lines.append("  …")
            elif k == "items":
                lines.append(f'  <span class="k">"items"</span>: [{{<span class="k">"sku"</span>: {js(r["items"][0]["sku"], 0, set())}, …}}, …]')
            else:
                lines.append(f'  <span class="k">"{k}"</span>: {js(r[k], 1, hl, k)}')
        return "{\n" + ",\n".join(lines) + "\n}"

    wa = src / "warehouse_receipts" / "CN01" / "NhapKho_CN01_20240313.xlsx"
    wb = src / "warehouse_receipts" / "CN03" / "receipts_CN03_20240904.xlsx"
    sh = src / "shopeefood" / "ShopeeFood_DoiSoat_20240311_20240317.xlsx"
    tk = src / "timekeeping" / "BR01" / "punch_BR01_20240311.csv"
    pos_new = src / "pos" / "POS-01" / "pos_export_2024-03-11.json"
    pos_old = src / "pos" / "POS-05" / "pos_export_2024-03-11.json"
    wb_rows = openpyxl.load_workbook(wb).active.max_row

    def is_odd_qty(v):
        m = re.match(r"([\d.]+)\s*kg$", str(v))
        return bool(m) and float(m.group(1)) > 100

    a = sheet(wa, 13, 6, more=True, hl=lambda r, c, v: r in (1, 2, 3, 11, 13))
    b = sheet(wb, wb_rows, 3, more=True, hl=lambda r, c, v: (c == 3 and r > 1) or (c == 2 and r > 1 and not re.search(r"[^\x00-\x7f]", str(v))))
    b_odd = [r for r in openpyxl.load_workbook(wb).active.iter_rows(min_row=2, values_only=True) if is_odd_qty(r[2])]
    s = sheet(sh, 8, 6, more=True, hl=lambda r, c, v: r in (1, 2) or (c == 1 and r > 4))
    with tk.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))[:6]
    t = "\n".join(",".join(f'<span class="hl">{esc(x)}</span>' if i in (0, 1) and j > 0 else esc(x)
                           for i, x in enumerate(row)) for j, row in enumerate(rows))
    bm = next(csv.DictReader((ROOT / "dbt" / "seeds" / "branch_code_map.csv").open(encoding="utf-8")))
    ids = [("POS", bm["pos_terminal_prefix"]), ("App riêng", bm["app_branch_code"]), ("GrabFood", bm["grab_store_name"]),
           ("ShopeeFood", bm["shopee_store_code"]), ("Chấm công", bm["timekeeping_branch_code"]), ("Kho, kiểm kê", bm["warehouse_folder"])]
    idt = "".join(f"<tr><td>{esc(k)}</td><td class='mono'>{esc(v)}</td></tr>" for k, v in ids)
    odd = "; ".join(f"{esc(r[1])} {esc(r[2])}" for r in b_odd)

    page = f"""<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="_style.css">
<style>
#fig {{ width: 790px; display: grid; grid-template-columns: 474px 1fr; gap: 12px; }}
.p {{ border: 1.3px solid #c5cdd5; border-radius: 9px; padding: 7px 9px 8px 9px; background: #fafbfc; min-width: 0; }}
.p h5 {{ margin: 0 0 5px 0; font-size: 13px; color: #1f2933; }} .p h5 span {{ font-weight: 400; color: #52606d; font-family: Consolas, monospace; font-size: 11.5px; }}
.p .note {{ font-size: 11.8px; color: #3e4c59; margin-top: 5px; line-height: 1.35; }}
.full {{ grid-column: 1 / -1; }}
.xl {{ border-collapse: collapse; font-family: Calibri, "Segoe UI", Arial, sans-serif; font-size: 12px; background: #fff; }}
.xl th {{ background: #eef1f4; color: #6b7785; font-weight: 400; border: 1px solid #d5dbe1; padding: 0 4px; font-size: 10.5px; }}
.xl td {{ border: 1px solid #e3e7ec; padding: 1px 5px; white-space: nowrap; max-width: 150px; overflow: hidden; text-overflow: ellipsis; }}
.xl td.num {{ text-align: right; }} .xl td.mg {{ text-align: center; font-weight: 700; }}
.hl {{ background: #fff1a8 !important; }}
pre {{ margin: 0; font-family: Consolas, "Cascadia Mono", monospace; font-size: 11.3px; line-height: 1.38; background: #fff;
       border: 1px solid #e3e7ec; padding: 5px 7px; white-space: pre; overflow: hidden; }}
.k {{ color: #7a3e9d; }} .s {{ color: #2f7d32; }} .n {{ color: #1d5fa8; }}
.two {{ display: grid; grid-template-columns: 1fr 1fr; gap: 10px; }} .two pre {{ white-space: pre-wrap; }}
.ids td {{ font-size: 12px; padding: 1px 10px 1px 0; white-space: nowrap; }} .ids td.mono {{ font-size: 11.8px; }}
</style></head><body><div id="fig">
<div class="p"><h5>(a) Sổ nhập kho mẫu A <span>{wa.name}</span></h5>{a}
  <div class="note">Tiêu đề phiếu, ô gộp, dòng Tổng cộng và chữ ký nằm lẫn trong bảng; tên nguyên liệu có dấu.</div></div>
<div class="p"><h5>(b) Sổ nhập kho mẫu B <span>{wb.name}</span></h5>{b}
  <div class="note">Số lượng dính đơn vị, tên không dấu hoặc viết khác; dòng {odd} là gam gõ vào cột kg.</div></div>
<div class="p full"><h5>(c) Báo cáo đối soát ShopeeFood <span>{sh.name}</span></h5>{s}
  <div class="note">Báo cáo theo tuần, hai dòng tiêu đề phía trên bảng; cửa hàng ghi theo hệ mã của sàn. Bảng gốc có 11 cột, hình chỉ hiện 6 cột đầu.</div></div>
<div class="p full"><h5>(d) Cùng một phần mềm POS, hai phiên bản <span>{pos_new.parent.name}/{pos_new.name} · {pos_old.parent.name}/{pos_old.name}</span></h5>
  <div class="two"><pre>{receipt(pos_new)}</pre><pre>{receipt(pos_old)}</pre></div>
  <div class="note">Khác định dạng ngày giờ, tiền là số hoặc chuỗi có dấu chấm ngăn cách, tên trường tổng tiền khác nhau. Hình lược bớt các trường không liên quan (…).</div></div>
<div class="p"><h5>(e) Máy chấm công <span>{tk.name}</span></h5><pre style="font-size:10.2px">{t}</pre>
  <div class="note">Mã nhân viên dạng số, không phải NV0xxx; mã chi nhánh BR01.</div></div>
<div class="p"><h5>(f) Một chi nhánh, sáu cách ghi mã</h5><table class="ids">{idt}</table>
  <div class="note">Tầng Silver ánh xạ về một mã chi nhánh duy nhất (seed <span class="mono">branch_code_map</span>).</div></div>
</div></body></html>"""
    (SRC / "h4_2_du_lieu_tho.html").write_text(page, encoding="utf-8")


# ---------------------------------------------------------------- xuat PNG
def render(names=None):
    gen_erd()
    if not names or "h4_2_du_lieu_tho" in names:
        gen_raw()
    files = sorted(p for p in SRC.iterdir() if p.suffix in (".dot", ".html") and not p.name.startswith("_"))
    if names:
        files = [p for p in files if p.stem in names]
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page(device_scale_factor=2, viewport={"width": 1100, "height": 800})
        for f in files:
            if f.suffix == ".dot":
                tmp = SRC / f"_tmp_{f.stem}.html"
                tmp.write_text(DOT_PAGE.format(dot=json.dumps(f.read_text(encoding="utf-8"))), encoding="utf-8")
                page.goto(tmp.as_uri())
                page.wait_for_selector("body[data-ready]")
                tmp.unlink()
            else:
                page.goto(f.as_uri())
                page.wait_for_load_state("networkidle")
            page.locator("#fig").screenshot(path=str(HERE / f"{f.stem}.png"))
            print("da dung", f.stem + ".png")
        browser.close()


if __name__ == "__main__":
    render(sys.argv[1:] or None)
